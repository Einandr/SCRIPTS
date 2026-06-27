import numpy as np
from openpyxl import load_workbook
import pandas as pd
from itertools import islice
from decimal import Decimal
import os

path = r'D:\YASIM\!VORON\!Properties'
file_name = 'properties_qubiq.xlsx'
sheet_name = 'steel_KBK-26'
material_name = 'steel'

os.chdir(path)
wb=load_workbook(file_name, data_only=True)

ws=wb[sheet_name]
print(wb.sheetnames)

df = pd.DataFrame()

for col in ws.iter_cols(min_row=0, values_only=True):
    df[col[0]]=col[1:]
    print(col)

df.dropna(axis=0, how='all', inplace=True)
# df.fillna(method='ffill', inplace=True)
# df.fillna(method='bfill', inplace=True)


# prop_density = ''.join(('/define/materials/change-create/', material_name, ' ', material_name, ' y piecewise-linear ', str(len(df.index))))
# for ind, item in enumerate(df['T_K']):
#     prop_density += ''.join((' ', str(df['T_K'][ind]), ' ', str(df['RO_kg-m3'][ind])))
# prop_density += ' n n n n n n\n'


def get_prop(str_start, x, y, quant, str_end):
    prop = str_start
    prop_points = ''
    n_points = 0
    for ind, item in enumerate(df[x]):
        if not np.isnan(df[y][ind]):
            n_points += 1
            prop_points += ''.join(
                (' ', str(df[x][ind]), ' ', str(Decimal(str(df[y][ind])).quantize(Decimal(quant)))))
    prop += str(n_points)
    prop += prop_points
    prop += str_end
    return prop


# str_start_Cp = ''.join(('/define/materials/change-create/', material_name, ' ', material_name, ' n y piecewise-linear '))
# str_end_Cp = ' n n n n n n\n'
# prop_Cp = get_prop(str_start_Cp, 'T_K', 'CP_J-kg-K', '.0001', str_end_Cp)

str_start_Lambda = ''.join(('/define/materials/change-create/', material_name, ' ', material_name, ' n n y piecewise-linear '))
str_end_Lambda = ' n n n n n n\n'
prop_Lambda = get_prop(str_start_Lambda, 'T [K]', 'Lam [Вт/м-К]', '.000001', str_end_Lambda)

# 'T_K', 'LAM_W-m-K'

# prop_viscocity = ''.join(('/define/materials/change-create/', material_name, ' ', material_name, ' n n n y piecewise-linear ', str(len(df.index))))
# for ind, item in enumerate(df['T_K']):
#     prop_viscocity += ''.join((' ', str(df['T_K'][ind]), ' ', str(Decimal(str(df['NU_Pa-s'][ind])).quantize(Decimal('.000001')))))
# prop_viscocity += ' n n n\n'

with open(''.join(('property_', material_name, '.jou')), 'w') as f:
    # f.write(';Density\n')
    # f.write(prop_density)
    # f.write(';Cp\n')
    # f.write(prop_Cp)
    f.write(';Lambda\n')
    f.write(prop_Lambda)
    # f.write(';Viscosity\n')
    # f.write(prop_viscocity)

print('the end')

