import os
from datetime import datetime

import pymorphy2
import comtypes.client

from pytrovich.detector import PetrovichGenderDetector
from pytrovich.enums import NamePart, Gender, Case
from pytrovich.maker import PetrovichDeclinationMaker

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import OxmlElement, parse_xml
from docx.shared import Inches, Pt, Cm, RGBColor


morph = pymorphy2.MorphAnalyzer()

# Словарь для месяцев на русском языке
months = {
    1: 'января',
    2: 'февраля',
    3: 'марта',
    4: 'апреля',
    5: 'мая',
    6: 'июня',
    7: 'июля',
    8: 'августа',
    9: 'сентября',
    10: 'октября',
    11: 'ноября',
    12: 'декабря'
}


def decline_word(number, one, few, many):
    """Склонение слов в зависимости от числа (например, 'рубль/рубля/рублей')"""
    number = abs(int(number)) % 100
    if 11 <= number <= 19:
        return many
    number = number % 10
    if number == 1:
        return one
    if 2 <= number <= 4:
        return few
    return many










def decline_fio_pytrovich(full_name):
    """Склонение ФИО в родительный падеж (например, 'Иванов Иван Иванович' → 'Иванова Ивана Ивановича')"""
    parts = full_name.split()
    if len(parts) != 3:
        return full_name

    surname, name, patronymic = parts
    print(surname, name, patronymic)
    maker = PetrovichDeclinationMaker()
    detector = PetrovichGenderDetector()
    gender = detector.detect(surname, name, patronymic)
    print(gender)
    declined_surname = maker.make(NamePart.LASTNAME, gender, Case.GENITIVE, surname)
    declined_name = maker.make(NamePart.FIRSTNAME, gender, Case.GENITIVE, name)
    declined_patronymic = maker.make(NamePart.MIDDLENAME, gender, Case.GENITIVE, patronymic)
    print(declined_surname, declined_name, declined_patronymic)
    return f"{declined_surname} {declined_name} {declined_patronymic}"

def decline_fio(full_name):
    """Склонение ФИО в родительный падеж (например, 'Иванов Иван Иванович' → 'Иванова Ивана Ивановича')"""
    parts = full_name.split()
    if len(parts) != 3:
        return full_name

    surname, name, patronymic = parts

    # Указываем конкретные теги: Surn (фамилия), Name (имя), Patr (отчество)
    def decline_part(word, tag):
        parsed = morph.parse(word)
        print(parsed, tag)
        # Ищем среди вариантов разбора тот, который совпадает по нужному типу
        best_match = next((p for p in parsed if tag in p.tag), parsed[0])
        declined = best_match.inflect({'gent'})
        return declined.word.capitalize() if declined else word.capitalize()

    declined_surname = decline_part(surname, 'Surn')
    declined_name = decline_part(name, 'Name')
    declined_patronymic = decline_part(patronymic, 'Patr')


    # # Склоняем фамилию (женские фамилии на -а/-я склоняются иначе)
    # parsed_surname = morph.parse(surname)[0]
    # declined_surname = parsed_surname.inflect({'gent'}).word.capitalize()
    #
    # # Склоняем имя
    # parsed_name = morph.parse(name)[0]
    # declined_name = parsed_name.inflect({'gent'}).word.capitalize()
    #
    # # Склоняем отчество
    # parsed_patronymic = morph.parse(patronymic)[0]
    # declined_patronymic = parsed_patronymic.inflect({'gent'}).word.capitalize()

    return f"{declined_surname} {declined_name} {declined_patronymic}"


def convert_docx_to_pdf(docx_path, pdf_path):
    """Конвертация DOCX в PDF с использованием Microsoft Word"""
    word = comtypes.client.CreateObject("Word.Application")
    doc = word.Documents.Open(docx_path)
    doc.SaveAs(pdf_path, FileFormat=17)  # 17 = wdFormatPDF
    doc.Close()
    word.Quit()


def num_to_words(num):
    """Преобразование числа в слова (например, 123 → 'сто двадцать три')"""
    units = ["", "один", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять"]
    teens = ["десять", "одиннадцать", "двенадцать", "тринадцать", "четырнадцать", "пятнадцать", "шестнадцать", "семнадцать", "восемнадцать", "девятнадцать"]
    tens = ["", "десять", "двадцать", "тридцать", "сорок", "пятьдесят", "шестьдесят", "семьдесят", "восемьдесят", "девяносто"]
    hundreds = ["", "сто", "двести", "триста", "четыреста", "пятьсот", "шестьсот", "семьсот", "восемьсот", "девятьсот"]
    thousands = ["", "тысяча", "тысячи", "тысяч"]

    def thousand(n):
        if n == 0:
            return ""
        res = []
        hundred = n // 100
        remainder = n % 100
        if hundred > 0:
            res.append(hundreds[hundred])
        if remainder >= 20:
            ten = remainder // 10
            unit = remainder % 10
            if ten > 0:
                res.append(tens[ten])
            if unit > 0:
                res.append(units[unit])
        elif remainder >= 10:
            res.append(teens[remainder - 10])
        elif remainder > 0:
            res.append(units[remainder])
        return " ".join(res)

    if num == 0:
        return "ноль"

    parts = []
    num = int(num)
    if num >= 1000:
        th = num // 1000
        parts.append(thousand(th))
        th = th % 10
        if th == 1:
            parts.append(thousands[1])
        elif 2 <= th <= 4:
            parts.append(thousands[2])
        else:
            parts.append(thousands[3])
        num = num % 1000
    if num > 0:
        parts.append(thousand(num))

    return " ".join(parts)


def setup_document_styles(doc, size_normal, size_heading_1, size_heading_2, line_spacing):
    """Настройка стилей документа (шрифт, отступы, выравнивание)"""
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(size_normal)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    style.paragraph_format.line_spacing = line_spacing

    # Настройка стиля "Заголовок 1"
    heading1_style = doc.styles['Heading 1']
    heading1_font = heading1_style.font
    heading1_font.name = 'Times New Roman'
    heading1_font.size = Pt(size_heading_1)
    heading1_font.color.rgb = RGBColor(0, 0, 0)
    heading1_style.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    heading1_style.paragraph_format.space_before = Pt(0)

    # Настройка стиля "Заголовок 2"
    heading2_style = doc.styles['Heading 2']
    heading2_font = heading2_style.font
    heading2_font.name = 'Times New Roman'
    heading2_font.size = Pt(size_heading_2)
    heading2_font.color.rgb = RGBColor(0, 0, 0)
    heading2_style.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    heading2_style.paragraph_format.space_before = Pt(5)


def add_text(paragraph, text, bold=False):
    """Добавление текста в абзац с возможностью выделения жирным"""
    run = paragraph.add_run(text)
    if bold:
        run.bold = True


def add_floating_image(paragraph, image_path, width, x_offset=0, y_offset=0):
    """Добавление изображения с обтеканием перед текстом"""
    run = paragraph.add_run()
    # 1. Сначала добавляем картинку как обычную (inline)
    img = run.add_picture(image_path, width=width)

    # 2. Получаем доступ к XML картинки
    # В новых версиях путь такой: img._inline
    inline = img._inline

    # 3. Генерируем XML для "плавающего" (anchor) размещения
    # simplePos="0" позволяет задавать точные координаты относительно текста
    anchor_xml = f'''
    <wp:anchor distT="0" distB="0" distL="114300" distR="114300" simplePos="0" relativeHeight="251658240" \
    behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1" {nsdecls('wp', 'r', 'a', 'pic', 'w')}>
        <wp:simplePos x="0" y="0"/>
        <wp:positionH relativeFrom="column">
            <wp:posOffset>{x_offset}</wp:posOffset> 
        </wp:positionH>
        <wp:positionV relativeFrom="line">
            <wp:posOffset>{y_offset}</wp:posOffset>
        </wp:positionV>
        <wp:extent cx="{inline.extent.cx}" cy="{inline.extent.cy}"/>
        <wp:effectExtent l="0" t="0" r="0" b="0"/>
        <wp:wrapNone/>
        <wp:docPr id="1" name="Signature"/>
        <wp:cNvGraphicFramePr>
            <a:graphicFrameLocks noChangeAspect="1"/>
        </wp:cNvGraphicFramePr>
        {inline.graphic.xml}
    </wp:anchor>
    '''

    # 4. Подменяем inline на anchor
    inline.getparent().replace(inline, parse_xml(anchor_xml))


